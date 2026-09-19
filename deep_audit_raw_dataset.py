import openpyxl
import pandas as pd
import numpy as np
import json
import sys
from collections import Counter

sys.stdout.reconfigure(encoding='utf-8')

excel_path = 'Paytm UPI Growth Challenge 2026 — Raw Short VOC Responses.xlsx'
wb = openpyxl.load_workbook(excel_path)
print("=== SHEETS IN WORKBOOK ===")
for s in wb.sheetnames:
    sheet = wb[s]
    print(f"Sheet '{s}': {sheet.max_row} rows, {sheet.max_column} cols")

df = pd.read_excel(excel_path, sheet_name='Form Responses 1')
print(f"\nTotal rows in 'Form Responses 1': {len(df)}")
print(f"Total columns: {len(df.columns)}")

print("\n=== COLUMN LIST ===")
for idx, col in enumerate(df.columns):
    print(f"[{idx+1}] {col}")

col_map = {
    'Timestamp': 'timestamp',
    '1. Which best describes you?': 'profile',
    '2. Which city do you currently live in?': 'city',
    '3. Which UPI apps do you currently use?': 'apps_used',
    '4. Which UPI app do you use MOST?': 'primary_app',
    '5. Have you used Paytm UPI in the last 90 days?': 'paytm_90d',
    '6. Roughly how many UPI payments do you make in a typical week?': 'weekly_freq',
    '7. Where do you use UPI MOST often?': 'top_occasions',
    '8. What is the BIGGEST reason you usually open your primary UPI app?': 'primary_reason',
    '9. When are you MOST likely to use Paytm UPI?': 'paytm_occasions',
    '10. What is the BIGGEST reason you do NOT use Paytm for more of your UPI payments?': 'paytm_barriers',
    '11. In ONE sentence: what would make you use Paytm more?': 'paytm_trigger_verbatim',
    '12. Would you be open to a 15-minute follow-up conversation about your UPI habits?': 'open_followup',
    '13. If yes, please share your preferred contact (phone/email).': 'contact'
}
df = df.rename(columns=col_map)

# Clean text fields
for c in df.columns:
    if df[c].dtype == 'object':
        df[c] = df[c].astype(str).str.replace('\ufffd', '–')

print("\n" + "="*80)
print("--- 1. DATA QUALITY & INTEGRITY AUDIT ---")
print("="*80)
print(f"Total Records: {len(df)}")
print(f"Unique Records: {len(df.drop_duplicates())}")
print(f"Duplicate records: {df.duplicated().sum()}")
print("Missing values per column:")
print(df.isnull().sum())

print("\n" + "="*80)
print("--- 2. GEOGRAPHIC DISTRIBUTION (REAL RAW DATA) ---")
print("="*80)
df['city_raw'] = df['city'].astype(str).str.strip()
df['city_clean'] = df['city_raw'].str.title().replace({
    'Bhopal ': 'Bhopal',
    'Uttar Pradesh ': 'Uttar Pradesh',
    'Up': 'Uttar Pradesh',
    'U.P': 'Uttar Pradesh',
    'U.P.': 'Uttar Pradesh',
    'Sitapur ': 'Sitapur',
    'Sitapur (U.P)': 'Sitapur',
    'Sitapur (U.P.)': 'Sitapur',
    'Lucknow ': 'Lucknow',
    'Lakhimpur': 'Lakhimpur Kheri',
    'Lakhimpur-Kheri': 'Lakhimpur Kheri',
    'Lakhimpur Kheri ': 'Lakhimpur Kheri'
})

print("City breakdown (Raw Counts & Percentages):")
city_summary = df['city_clean'].value_counts()
city_pct = (city_summary / len(df) * 100).round(2)
print(pd.DataFrame({'Count': city_summary, 'Percentage': city_pct}))

# State / Region mapping based on actual cities in the dataset
def map_state(city):
    c = str(city).lower()
    if any(x in c for x in ['sitapur', 'lucknow', 'lakhimpur', 'shahjahanpur', 'biswan', 'hardoi', 'ballia', 'uttar pradesh', 'agra', 'noida', 'babhnan', 'up']):
        return 'Uttar Pradesh'
    elif any(x in c for x in ['bhopal', 'sehore', 'mp', 'madhya pradesh']):
        return 'Madhya Pradesh'
    elif any(x in c for x in ['bangalore', 'bengaluru', 'karnataka']):
        return 'Karnataka'
    elif any(x in c for x in ['delhi', 'new delhi', 'ncr']):
        return 'Delhi (NCR)'
    elif any(x in c for x in ['dehradun', 'uttarakhand']):
        return 'Uttarakhand'
    elif any(x in c for x in ['panipat', 'haryana', 'gurgaon', 'gurugram']):
        return 'Haryana'
    elif any(x in c for x in ['shimla', 'himachal']):
        return 'Himachal Pradesh'
    elif any(x in c for x in ['surat', 'gujarat']):
        return 'Gujarat'
    else:
        return 'Other / Unspecified'

df['state_mapped'] = df['city_clean'].apply(map_state)
print("\nState / Region breakdown (from real dataset):")
state_summary = df['state_mapped'].value_counts()
state_pct = (state_summary / len(df) * 100).round(2)
print(pd.DataFrame({'Count': state_summary, 'Percentage': state_pct}))

# Urban vs Rural mapping based on actual locations in dataset
# Tier 1/2 Metros vs Tier 3 / Semi-Urban / Rural Towns in UP/MP
def map_urban_rural(city):
    c = str(city).lower()
    if any(x in c for x in ['bangalore', 'bengaluru', 'delhi', 'noida', 'gurgaon', 'surat', 'lucknow', 'bhopal', 'dehradun']):
        return 'Urban / Metro / Tier 1-2'
    elif any(x in c for x in ['sitapur', 'lakhimpur', 'biswan', 'hardoi', 'ballia', 'shahjahanpur', 'sehore', 'babhnan', 'panipat', 'shimla']):
        return 'Semi-Urban / Rural / Tier 3 Towns'
    else:
        return 'Semi-Urban / Tier 3'

df['urban_rural_mapped'] = df['city_clean'].apply(map_urban_rural)
print("\nUrban vs Semi-Urban / Tier-3 Town Classification (from real dataset):")
ur_summary = df['urban_rural_mapped'].value_counts()
ur_pct = (ur_summary / len(df) * 100).round(2)
print(pd.DataFrame({'Count': ur_summary, 'Percentage': ur_pct}))

print("\n" + "="*80)
print("--- 3. RESPONDENT PROFILE & OCCUPATION (REAL RAW DATA) ---")
print("="*80)
prof_summary = df['profile'].value_counts()
prof_pct = (prof_summary / len(df) * 100).round(2)
print(pd.DataFrame({'Count': prof_summary, 'Percentage': prof_pct}))

print("\n" + "="*80)
print("--- 4. PRIMARY UPI APP SHARE (REAL RAW DATA) ---")
print("="*80)
df['primary_app_clean'] = df['primary_app'].replace({'Google pay': 'Google Pay'}).str.strip()
app_summary = df['primary_app_clean'].value_counts()
app_pct = (app_summary / len(df) * 100).round(2)
print(pd.DataFrame({'Count': app_summary, 'Percentage': app_pct}))

print("\n" + "="*80)
print("--- 5. PAYTM 90-DAY ACTIVITY (REAL RAW DATA) ---")
print("="*80)
p90_summary = df['paytm_90d'].value_counts()
p90_pct = (p90_summary / len(df) * 100).round(2)
print(pd.DataFrame({'Count': p90_summary, 'Percentage': p90_pct}))

print("\n" + "="*80)
print("--- 6. WEEKLY TRANSACTION FREQUENCY (REAL RAW DATA) ---")
print("="*80)
wf_summary = df['weekly_freq'].value_counts()
wf_pct = (wf_summary / len(df) * 100).round(2)
print(pd.DataFrame({'Count': wf_summary, 'Percentage': wf_pct}))

print("\n" + "="*80)
print("--- 7. DEEP CROSS-TABULATION: PRIMARY APP vs PAYTM 90D USAGE ---")
print("="*80)
ct_app_p90 = pd.crosstab(df['primary_app_clean'], df['paytm_90d'], margins=True)
print(ct_app_p90)
print("\nRow Percentages:")
print((pd.crosstab(df['primary_app_clean'], df['paytm_90d'], normalize='index') * 100).round(2))

print("\n" + "="*80)
print("--- 8. DEEP CROSS-TABULATION: PROFILE vs PRIMARY APP ---")
print("="*80)
ct_prof_app = pd.crosstab(df['profile'], df['primary_app_clean'], margins=True)
print(ct_prof_app)
print("\nRow Percentages:")
print((pd.crosstab(df['profile'], df['primary_app_clean'], normalize='index') * 100).round(2))

print("\n" + "="*80)
print("--- 9. MULTI-SELECT FIELDS DETAILED FREQUENCY ANALYSIS ---")
print("="*80)

def analyze_multiselect(series, name):
    items = []
    for val in series:
        if pd.isna(val) or val in ['None', 'nan', '']:
            continue
        parts = [p.strip() for p in str(val).split(',') if p.strip()]
        items.extend(parts)
    c = Counter(items)
    print(f"\n>>> {name} (Total Mentions: {len(items)}, Unique Items: {len(c)})")
    res = []
    for k, v in c.most_common():
        pct = (v / len(series)) * 100
        res.append({'Item': k, 'Mentions': v, '% of Respondents (N=111)': round(pct, 2)})
    print(pd.DataFrame(res).to_string(index=False))
    return res

apps_repertoire_res = analyze_multiselect(df['apps_used'], "APPS IN REPERTOIRE (Q3)")
top_occasions_res = analyze_multiselect(df['top_occasions'], "TOP GENERAL UPI OCCASIONS (Q7)")
primary_reasons_res = analyze_multiselect(df['primary_reason'], "REASONS FOR CHOOSING PRIMARY APP (Q8)")
paytm_occasions_res = analyze_multiselect(df['paytm_occasions'], "WHEN USERS USE PAYTM UPI (Q9)")
paytm_barriers_res = analyze_multiselect(df['paytm_barriers'], "PAYTM BARRIERS (Q10)")

print("\n" + "="*80)
print("--- 10. VERBATIMS IN-DEPTH THEMATIC & SEMANTIC ANALYSIS ---")
print("="*80)

# Export all processed data to a full json audit file
df.to_json('real_dataset_full_audit.json', orient='records', indent=2, force_ascii=False)
print("Saved real_dataset_full_audit.json successfully.")
